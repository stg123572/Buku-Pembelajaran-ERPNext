import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_execution_guide(docx_path):
    doc = Document()

    # Set Margins 1 inch
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # Styles
    normal = doc.styles['Normal']
    normal.font.name = 'Segoe UI'
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(44, 62, 80)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.bold = True
        r.font.name = 'Segoe UI'
        r.font.size = Pt(15)
        r.font.color.rgb = RGBColor(27, 79, 114)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(13)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.bold = True
        r.font.name = 'Segoe UI'
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor(41, 128, 185)
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(9)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.bold = True
        r.font.name = 'Segoe UI'
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(52, 73, 94)
        return p

    def add_p(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.bold = True
        p.add_run(text)
        return p

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.bold = True
        p.add_run(text)
        return p

    def add_num(num_str, bold_prefix, text=""):
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.space_before = Pt(1.5)
        p.paragraph_format.space_after = Pt(2)
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.bold = True
        if text:
            p.add_run(text)
        return p

    def add_callout(text, title="TIPS PRAKTIK:"):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.left_indent = Inches(0.25)
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="8" w:color="2E86C1"/></w:pBdr>')
        p._p.get_or_add_pPr().append(pBdr)
        r_title = p.add_run(f"[{title}] ")
        r_title.bold = True
        r_title.font.color.rgb = RGBColor(46, 134, 193)
        p.add_run(text)

    def add_styled_table(headers, rows):
        t = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(t)

        # Header
        for ci, h in enumerate(headers):
            cell = t.rows[0].cells[ci]
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            set_cell_background(cell, "1B4F72")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(h)
            r.bold = True
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(255, 255, 255)

        # Rows
        for ri, rdata in enumerate(rows):
            for ci, val in enumerate(rdata):
                cell = t.rows[ri + 1].cells[ci]
                set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                bg = "F8F9FA" if ri % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(1.5)
                p.paragraph_format.space_after = Pt(1.5)
                r = p.add_run(val)
                r.font.size = Pt(8.5)

    # HEADER & TITLE
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("PANDUAN PELAKSANAAN PRAKTIK ERPNEXT\nPENYELESAIAN 5 SOAL STUDI KASUS MANUFAKTUR")
    r_title.bold = True
    r_title.font.name = 'Segoe UI'
    r_title.font.size = Pt(16)
    r_title.font.color.rgb = RGBColor(27, 79, 114)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("Panduan Klik-demi-Klik (GUI Desk) Modul Stock, Manufacturing, Buying, Selling, & Accounting\nPerusahaan: PT Anggi Karya Manufaktur (AnggiST)")
    r_sub.font.name = 'Segoe UI'
    r_sub.font.size = Pt(10)
    r_sub.font.color.rgb = RGBColor(90, 107, 124)

    # Divider
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(10)
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="1F4E79"/></w:pBdr>')
    p_div._p.get_or_add_pPr().append(pBdr)

    # PETUNJUK AWAL
    add_h1("Petunjuk Awal: Persiapan Lingkungan ERPNext")
    add_bullet("URL Akses Sistem: ", "Buka browser Google Chrome / Microsoft Edge di alamat http://localhost:8080")
    add_bullet("Akun Login: ", "Username: Administrator | Password: admin (atau akun Anda: anggisitanggang056@gmail.com)")
    add_bullet("Perusahaan Aktif: ", "Pastikan perusahaan yang dipilih adalah AnggiST (Singkatan Gudang: - A).")

    # ================= MODUL 1: SOAL 1 =================
    add_h1("MODUL 1 (SOAL 1): Konfigurasi Master Data & Saldo Awal")
    add_p("Tujuan: Membangun fondasi master klasifikasi barang, hierarki gudang, mitra bisnis, master item lengkap, dan mengisi 50 unit bahan baku sebagai saldo awal persediaan.")

    add_h2("Langkah 1.1: Membuat Kelompok Barang (Item Group)")
    add_num("1", "Navigasi: ", "Ketik di search bar atas 'Item Group List' atau klik Stock > Item Group.")
    add_num("2", "Aksi: ", "Klik tombol '+ Add Item Group' di pojok kanan atas.")
    add_num("3", "Pengisian Form: ", "Isi 'Item Group Name': Bahan Baku, 'Parent Item Group': All Item Groups, 'Group Node': Jangan dicentang. Klik Save.")
    add_num("4", "Ulangi: ", "Buat kelompok lainnya dengan cara yang sama: Barang Jadi, Bahan Kemasan, Barang Setengah Jadi, dan Sub-Assembly.")

    add_h2("Langkah 1.2: Membuat Master Gudang (Warehouse)")
    add_num("1", "Navigasi: ", "Buka Stock > Warehouse (atau search 'Warehouse List').")
    add_num("2", "Aksi: ", "Klik '+ Add Warehouse'.")
    add_num("3", "Pengisian Form: ", "Tambahkan 4 gudang berikut satu per satu:")
    add_bullet("Gudang 1: ", "Warehouse Name: 'Gudang Bahan Baku', Company: 'AnggiST', Parent Warehouse: 'All Warehouses - A'. Klik Save (Sistem otomatis memberi ID: Gudang Bahan Baku - A).")
    add_bullet("Gudang 2: ", "Warehouse Name: 'Gudang Work In Progress', Company: 'AnggiST', Parent Warehouse: 'All Warehouses - A'.")
    add_bullet("Gudang 3: ", "Warehouse Name: 'Gudang Barang Jadi', Company: 'AnggiST', Parent Warehouse: 'All Warehouses - A'.")
    add_bullet("Gudang 4: ", "Warehouse Name: 'Gudang Reject dan Scrap', Company: 'AnggiST', Parent Warehouse: 'All Warehouses - A'.")

    add_h2("Langkah 1.3: Membuat Pemasok (Supplier) & Pelanggan (Customer)")
    add_bullet("Supplier: ", "Buying > Supplier > '+ Add Supplier'. Masukkan Supplier Name: 'PT Baja Mandiri Perkasa', Supplier Group: 'Hardware'. Ulangi untuk CV Foam & Textile Nusantara, PT Mitra Hardware Indonesia, CV Karton Box Lestari.")
    add_bullet("Customer: ", "Selling > Customer > '+ Add Customer'. Masukkan Customer Name: 'PT Sejahtera Finansial Utama', Customer Group: 'Commercial', Territory: 'Indonesia'. Ulangi untuk Toko Mebel Modern Jaya, CV Digital Solusi Mandiri.")

    add_h2("Langkah 1.4: Membuat Master Items (Produk Jadi & Bahan Baku)")
    add_num("1", "Navigasi: ", "Stock > Item > '+ Add Item'.")
    add_num("2", "Input Barang Jadi: ", "Item Code: 'FG-KURSI-01', Item Name: 'Kursi Kantor Ergonomis Nexus', Item Group: 'Barang Jadi', Default UOM: 'Nos', Standard Selling Rate: 1.250.000, Valuation Rate: 650.000, Default Warehouse: 'Gudang Barang Jadi - A'. Klik Save.")
    add_num("3", "Input 10 Komponen Bahan Baku & Kemasan: ", "Gunakan rincian data pada tabel berikut:")

    item_rows = [
        ["RM-RANGKA-01", "Rangka Punggung Besi", "Bahan Baku", "Nos", "150.000", "Gudang Bahan Baku - A"],
        ["RM-BUSA-01", "Busa Dudukan High Density", "Bahan Baku", "Nos", "75.000", "Gudang Bahan Baku - A"],
        ["RM-KAIN-01", "Kain Pelapis Mesh Hitam", "Bahan Baku", "Meter", "45.000", "Gudang Bahan Baku - A"],
        ["RM-GASLIFT-01", "Silinder Hidrolik Gas Lift Class 3", "Bahan Baku", "Nos", "85.000", "Gudang Bahan Baku - A"],
        ["RM-BASE-01", "Kaki Bintang Lima Nilon 320mm", "Bahan Baku", "Nos", "95.000", "Gudang Bahan Baku - A"],
        ["RM-RODA-01", "Set Roda Caster 50mm (5 Pcs)", "Bahan Baku", "Set", "40.000", "Gudang Bahan Baku - A"],
        ["RM-ARMREST-01", "Set Sandaran Tangan 3D", "Bahan Baku", "Set", "110.000", "Gudang Bahan Baku - A"],
        ["RM-BAUT-01", "Paket Baut & Kunci L", "Bahan Baku", "Set", "20.000", "Gudang Bahan Baku - A"],
        ["PKG-BOX-01", "Kardus Kemasan Corrugated Brown", "Bahan Kemasan", "Nos", "25.000", "Gudang Bahan Baku - A"],
        ["PKG-PLASTIK-01", "Plastik Pembungkus & Bubble Wrap", "Bahan Kemasan", "Roll", "10.000", "Gudang Bahan Baku - A"],
    ]
    add_styled_table(["Kode Item", "Nama Item", "Kelompok", "Satuan", "Harga Nilai (Rp)", "Gudang Default"], item_rows)

    add_h2("Langkah 1.5: Input Saldo Awal Persediaan (Material Receipt)")
    add_num("1", "Navigasi: ", "Stock > Stock Entry > '+ Add Stock Entry'.")
    add_num("2", "Header Dokumen: ", "Pilih Stock Entry Type: 'Material Receipt' (Purpose otomatis menjadi Material Receipt).")
    add_num("3", "Default Warehouse: ", "Pilih Default Target Warehouse: 'Gudang Bahan Baku - A'.")
    add_num("4", "Tabel Items: ", "Tambahkan 10 baris item di atas dengan Qty masing-masing 50 (khusus RM-KAIN-01 isi 100 Meter) dan Basic Rate sesuai harga standar masing-masing.")
    add_num("5", "Akun Penyeimbang: ", "Pastikan Expense Account terisi '5110.020 - Penyesuaian Stock - A'.")
    add_num("6", "Simpan & Submit: ", "Klik 'Save', lalu klik tombol hijau 'Submit' di pojok kanan atas. Status berubah menjadi 'Submitted'.")
    add_callout("Total nilai persediaan awal yang terbentuk di Gudang Bahan Baku adalah Rp 35.000.000. Verifikasi melalui menu Stock > Stock Balance.", "VERIFIKASI SOAL 1")

    # ================= MODUL 2: SOAL 2 =================
    add_h1("MODUL 2 (SOAL 2): Setup Workstation, Operation & BOM")
    add_p("Tujuan: Mengonfigurasi stasiun kerja lantai pabrik, jenis operasi perakitan, dan resep produksi (Bill of Materials) 1 unit kursi kantor ergonomis.")

    add_h2("Langkah 2.1: Menyiapkan Workstations (Stasiun Kerja)")
    add_num("1", "Navigasi: ", "Manufacturing > Workstation > '+ Add Workstation'.")
    add_num("2", "Input Stasiun 1: ", "Workstation Name: 'WS Potong & Jahit', Hour Rate: 45.000. Klik Save.")
    add_num("3", "Input Stasiun 2: ", "Workstation Name: 'WS Perakitan Utama', Hour Rate: 50.000. Klik Save.")
    add_num("4", "Input Stasiun 3: ", "Workstation Name: 'WS QC & Packaging', Hour Rate: 35.000. Klik Save.")

    add_h2("Langkah 2.2: Menyiapkan Operations (Operasi Produksi)")
    add_num("1", "Navigasi: ", "Manufacturing > Operation > '+ Add Operation'.")
    add_num("2", "Operasi 1: ", "Operation Name: 'OP-JAHIT-BUSA', Default Workstation: 'WS Potong & Jahit'. Klik Save.")
    add_num("3", "Operasi 2: ", "Operation Name: 'OP-ASSEMBLY-FRAME', Default Workstation: 'WS Perakitan Utama'. Klik Save.")
    add_num("4", "Operasi 3: ", "Operation Name: 'OP-QC-PACKING', Default Workstation: 'WS QC & Packaging'. Klik Save.")

    add_h2("Langkah 2.3: Membuat Formula Bill of Materials (BOM)")
    add_num("1", "Navigasi: ", "Manufacturing > Bill of Materials > '+ Add BOM'.")
    add_num("2", "Header BOM: ", "Item: 'FG-KURSI-01', Quantity: 1, Currency: IDR, centang 'With Operations'.")
    add_num("3", "Tabel Operations (Routing): ", "Tambahkan 3 operasi berurutan:")
    add_bullet("Baris 1: ", "Operation: OP-JAHIT-BUSA | Workstation: WS Potong & Jahit | Time in Mins: 30")
    add_bullet("Baris 2: ", "Operation: OP-ASSEMBLY-FRAME | Workstation: WS Perakitan Utama | Time in Mins: 20")
    add_bullet("Baris 3: ", "Operation: OP-QC-PACKING | Workstation: WS QC & Packaging | Time in Mins: 10")
    add_num("4", "Tabel Materials (Bahan Baku): ", "Masukkan 10 komponen (1 Rangka, 1 Busa, 2 Meter Kain, 1 Gaslift, 1 Kaki Bintang, 1 Set Roda, 1 Set Armrest, 1 Set Baut, 1 Kardus, 1 Rol Plastik).")
    add_num("5", "Simpan & Submit: ", "Klik 'Save', lalu klik 'Submit'.")
    add_callout("Sistem otomatis mengkalkulasi HPP Standar Kursi: Biaya Material (Rp 700.000) + Biaya Operasi (Rp 45.000) = Rp 745.000 / unit. Dokumen tersimpan dengan ID BOM-FG-KURSI-01-001.", "VERIFIKASI SOAL 2")

    # ================= MODUL 3: SOAL 3 =================
    add_h1("MODUL 3 (SOAL 3): Siklus Pengadaan / Procurement Flow")
    add_p("Tujuan: Mempraktikkan alur pengadaan barang mulai dari Purchase Order -> Purchase Receipt -> Purchase Invoice.")

    add_h2("Langkah 3.1: Menerbitkan Purchase Order (PO-SIM-001)")
    add_num("1", "Navigasi: ", "Buying > Purchase Order > '+ Add Purchase Order'.")
    add_num("2", "Pengisian Data: ", "Supplier: 'PT Baja Mandiri Perkasa', Required By Date: hari ini.")
    add_num("3", "Tabel Items: ", "Tambahkan 2 baris item pengadaan:")
    add_bullet("Item 1: ", "RM-RANGKA-01, Qty: 20, Rate: 150.000, Warehouse: 'Gudang Bahan Baku - A'.")
    add_bullet("Item 2: ", "RM-BASE-01, Qty: 20, Rate: 95.000, Warehouse: 'Gudang Bahan Baku - A'.")
    add_num("4", "Simpan & Submit: ", "Klik 'Save' lalu klik 'Submit'. Total PO: Rp 4.900.000. Dokumen: PUR-ORD-2026-00011.")

    add_h2("Langkah 3.2: Penerimaan Barang Fisik (Purchase Receipt)")
    add_num("1", "Aksi Cepat: ", "Pada halaman Purchase Order yang baru disubmit, klik tombol 'Create ▾' di kanan atas > pilih 'Receipt'.")
    add_num("2", "Verifikasi Form: ", "Sistem otomatis mengisi formulir Purchase Receipt dengan item dan kuantitas dari PO. Pastikan Accepted Warehouse adalah 'Gudang Bahan Baku - A'.")
    add_num("3", "Simpan & Submit: ", "Klik 'Save' lalu klik 'Submit'. Kuantitas fisik bertambah di gudang (+20 Rangka & +20 Kaki Bintang).")

    add_h2("Langkah 3.3: Pencatatan Tagihan Vendor (Purchase Invoice)")
    add_num("1", "Aksi Cepat: ", "Pada dokumen Purchase Receipt tadi, klik tombol 'Create ▾' > pilih 'Invoice'.")
    add_num("2", "Pengisian Faktur: ", "Isi Bill No (No Faktur Vendor): 'INV-BMP-2026/09'. Pastikan Credit To terisi '2111.001 - Hutang Dagang Dalam Negeri - A'.")
    add_num("3", "Simpan & Submit: ", "Klik 'Save' lalu klik 'Submit'. Hutang dagang resmi diakui di neraca sebesar Rp 4.900.000.")
    add_callout("Cek Accounting > General Ledger. Transaksi ini mendebit Persediaan Barang Rp 4.900.000 dan mengkredit Hutang Dagang Rp 4.900.000.", "VERIFIKASI SOAL 3")

    # ================= MODUL 4: SOAL 4 =================
    add_h1("MODUL 4 (SOAL 4): Eksekusi Manufaktur & Mutasi Gudang")
    add_p("Tujuan: Menjalankan perintah kerja perakitan 10 unit kursi, mentransfer bahan ke WIP, menyelesaikan Job Cards stasiun kerja, dan membukukan barang jadi.")

    add_h2("Langkah 4.1: Menerbitkan Work Order (Surat Perintah Kerja)")
    add_num("1", "Navigasi: ", "Manufacturing > Work Order > '+ Add Work Order'.")
    add_num("2", "Pengisian Form: ", "Item to Manufacture: 'FG-KURSI-01', Qty to Manufacture: 10, BOM No: otomatis terisi BOM-FG-KURSI-01-001.")
    add_num("3", "Pemilihan Gudang: ", "Source Warehouse: 'Gudang Bahan Baku - A', WIP Warehouse: 'Gudang Work In Progress - A', FG Warehouse: 'Gudang Barang Jadi - A'.")
    add_num("4", "Simpan & Submit: ", "Klik 'Save' lalu klik 'Submit'. Dokumen: MFG-WO-2026-00001 (Status: Not Started).")

    add_h2("Langkah 4.2: Transfer Bahan Baku ke Lantai Produksi (WIP Transfer)")
    add_num("1", "Aksi Cepat: ", "Di dalam halaman Work Order yang sudah disubmit, klik tombol 'Start' (atau klik 'Create ▾' > 'Transfer').")
    add_num("2", "Form Stock Entry: ", "Sistem membuat Stock Entry tipe 'Material Transfer for Manufacture'. Periksa bahwa 10 set bahan berpindah dari Gudang Bahan Baku ke Gudang WIP.")
    add_num("3", "Simpan & Submit: ", "Klik 'Save' lalu klik 'Submit'. Status Work Order berubah menjadi 'In Process'.")

    add_h2("Langkah 4.3: Menyelesaikan Pengerjaan Job Cards Stasiun Kerja")
    add_num("1", "Navigasi: ", "Buka Manufacturing > Job Card. Terlihat 3 Job Card yang dihasilkan dari Work Order.")
    add_num("2", "Selesaikan Berurutan: ", "Pastikan menyelesaikan sesuai nomor sequence:")
    add_bullet("Urutan 1: ", "Buka Job Card 'OP-JAHIT-BUSA' > Masukkan Time Log (30 menit, Qty: 10) > Klik 'Save' & 'Submit'.")
    add_bullet("Urutan 2: ", "Buka Job Card 'OP-ASSEMBLY-FRAME' > Masukkan Time Log (20 menit, Qty: 10) > Klik 'Save' & 'Submit'.")
    add_bullet("Urutan 3: ", "Buka Job Card 'OP-QC-PACKING' > Masukkan Time Log (10 menit, Qty: 10) > Klik 'Save' & 'Submit'.")

    add_h2("Langkah 4.4: Menyelesaikan Produksi (Manufacture / Finish)")
    add_num("1", "Aksi Cepat: ", "Kembali ke dokumen Work Order, klik tombol 'Finish' (atau 'Create ▾' > 'Manufacture').")
    add_num("2", "Verifikasi Form: ", "Sistem membuat Stock Entry tipe 'Manufacture' dengan 10 item komponen terpotong dari Gudang WIP dan 10 Unit FG-KURSI-01 masuk ke Gudang Barang Jadi.")
    add_num("3", "Simpan & Submit: ", "Klik 'Save' lalu klik 'Submit'. Status Work Order resmi menjadi 'COMPLETED'.")
    add_callout("Di Gudang Barang Jadi kini terdapat 10 unit Kursi Ergonomis dengan HPP Realisasi Rp 765.000/unit (Total Rp 7.650.000). Saldo bahan di Gudang WIP kembali 0.", "VERIFIKASI SOAL 4")

    # ================= MODUL 5: SOAL 5 =================
    add_h1("MODUL 5 (SOAL 5): Penjualan & Analisis Laba Kotor")
    add_p("Tujuan: Memproses order penjualan 10 unit kursi, mengirim barang fisik via surat jalan, menerbitkan faktur tagihan, dan mengevaluasi laba kotor.")

    add_h2("Langkah 5.1: Menerbitkan Sales Order (SO-SIM-001)")
    add_num("1", "Navigasi: ", "Selling > Sales Order > '+ Add Sales Order'.")
    add_num("2", "Pengisian Form: ", "Customer: 'PT Sejahtera Finansial Utama', Delivery Date: hari ini.")
    add_num("3", "Tabel Items: ", "Item Code: 'FG-KURSI-01', Qty: 10, Rate: 1.250.000, Warehouse: 'Gudang Barang Jadi - A'. Total: Rp 12.500.000.")
    add_num("4", "Simpan & Submit: ", "Klik 'Save' lalu klik 'Submit'. Dokumen: SAL-ORD-2026-00006.")

    add_h2("Langkah 5.2: Pengiriman Barang (Delivery Note / Surat Jalan)")
    add_num("1", "Aksi Cepat: ", "Pada Sales Order yang baru disubmit, klik tombol 'Create ▾' di kanan atas > pilih 'Delivery'.")
    add_num("2", "Verifikasi Form: ", "Pastikan Warehouse asal adalah 'Gudang Barang Jadi - A' dan kuantitas adalah 10 unit.")
    add_num("3", "Simpan & Submit: ", "Klik 'Save' lalu klik 'Submit'. Stok fisik kursi di Gudang Barang Jadi berkurang menjadi 0.")

    add_h2("Langkah 5.3: Penerbitan Faktur Penjualan (Sales Invoice)")
    add_num("1", "Aksi Cepat: ", "Pada Delivery Note tadi, klik 'Create ▾' > pilih 'Invoice'.")
    add_num("2", "Verifikasi Akun: ", "Debit To: '1131.0010 - Piutang Dagang - A', Income Account: '4110.000 - Penjualan - A'. Total Faktur: Rp 12.500.000.")
    add_num("3", "Simpan & Submit: ", "Klik 'Save' lalu klik 'Submit'. Piutang dagang dan pendapatan penjualan resmi tercatat.")

    add_h2("Langkah 5.4: Mengevaluasi Laporan Laba Kotor (Gross Profit)")
    add_num("1", "Navigasi Laba Rugi: ", "Buka Accounting > Profit and Loss Statement.")
    add_num("2", "Hasil Analisis: ", "Perhatikan struktur angka keuangan:")
    add_bullet("Pendapatan Penjualan (Revenue): ", "Rp 12.500.000,00 (10 Unit @ Rp 1.250.000)")
    add_bullet("Beban Pokok Penjualan (COGS / HPP): ", "Rp 7.650.000,00 (10 Unit @ Rp 765.000)")
    add_bullet("LABA KOTOR (GROSS PROFIT): ", "Rp 4.850.000,00")
    add_bullet("MARGIN LABA KOTOR (%): ", "38,80%")
    add_callout("Laba kotor sebesar Rp 4.850.000 (margin 38,80%) membuktikan efisiensi proses perakitan kursi kantor ergonomis pada PT Anggi Karya Manufaktur.", "VERIFIKASI SOAL 5")

    doc.save(docx_path)
    print("[SUCCESS] Dokumen Panduan Pelaksanaan DOCX berhasil dibuat di:", docx_path)

if __name__ == '__main__':
    doc_out = "/tmp/PANDUAN_PELAKSANAAN_PRAKTIK_ERPNEXT.docx"
    build_execution_guide(doc_out)
